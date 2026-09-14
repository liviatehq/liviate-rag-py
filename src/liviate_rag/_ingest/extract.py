"""Plain-text extraction per supported source_type.

Each function takes raw bytes and returns UTF-8 text ready for chunking.
Markdown/plain text are returned as-is (their own syntax is already
readable, no need to strip it for a RAG chunk); the others are converted
to plain text since their native format has no value once chunked.
"""

from __future__ import annotations

import csv
import io
import json

from ..exceptions import UnsupportedFileType


def extract_text(data: bytes, source_type: str) -> str:
    if source_type in ("md", "txt"):
        return data.decode("utf-8", errors="replace")
    if source_type == "pdf":
        return _extract_pdf(data)
    if source_type == "docx":
        return _extract_docx(data)
    if source_type == "html":
        return _extract_html(data)
    if source_type == "csv":
        return _extract_csv(data)
    if source_type == "json":
        return _extract_json(data)
    raise UnsupportedFileType(f"No text extractor for source_type={source_type!r}.")


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages).strip()


def _extract_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n\n".join(parts).strip()


class _HTMLTextExtractor:
    """Minimal tag-stripping HTML->text using the stdlib parser -- no extra
    dependency for something this simple. Drops <script>/<style> content
    entirely rather than including their non-prose text."""

    _SKIP_TAGS = {"script", "style", "noscript"}

    def __init__(self) -> None:
        from html.parser import HTMLParser

        self._chunks: list[str] = []
        self._skip_depth = 0

        outer = self

        class _Parser(HTMLParser):
            def handle_starttag(self, tag: str, attrs: list) -> None:  # noqa: ARG002
                if tag in outer._SKIP_TAGS:
                    outer._skip_depth += 1
                elif tag in ("p", "br", "div", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"):
                    outer._chunks.append("\n")

            def handle_endtag(self, tag: str) -> None:
                if tag in outer._SKIP_TAGS and outer._skip_depth > 0:
                    outer._skip_depth -= 1

            def handle_data(self, data: str) -> None:
                if outer._skip_depth == 0 and data.strip():
                    outer._chunks.append(data.strip())

        self._parser = _Parser()

    def feed(self, html: str) -> str:
        self._parser.feed(html)
        text = " ".join(self._chunks)
        return "\n".join(line.strip() for line in text.split("\n") if line.strip())


def _extract_html(data: bytes) -> str:
    return _HTMLTextExtractor().feed(data.decode("utf-8", errors="replace"))


def _extract_csv(data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return ""
    header, body = rows[0], rows[1:]
    lines = []
    for row in body:
        pairs = [f"{h}: {v}" for h, v in zip(header, row)]
        lines.append(", ".join(pairs))
    return "\n".join(lines)


def _extract_json(data: bytes) -> str:
    parsed = json.loads(data.decode("utf-8", errors="replace"))
    return json.dumps(parsed, indent=2, ensure_ascii=False)
